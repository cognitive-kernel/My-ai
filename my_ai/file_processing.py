from __future__ import annotations

import hashlib
import importlib.util
import mimetypes
import os
import platform
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

OPTIONAL_PREREQUISITES: dict[str, str] = {
    "docx": "python-docx",
    "openpyxl": "openpyxl",
    "pptx": "python-pptx",
    "fitz": "PyMuPDF",
    "PIL": "Pillow",
}

SYSTEM_PREREQUISITES = {"ffmpeg": "ffmpeg", "ffprobe": "ffprobe"}


def detect_type(path: str) -> dict[str, str]:
    p = Path(path)
    mime = mimetypes.guess_type(p.name)[0] or "application/octet-stream"
    suffix = p.suffix.lower()
    kind = "unknown"
    if mime.startswith("image/") or suffix in {".heic", ".heif"}:
        kind = "image"
    elif mime.startswith("audio/") or suffix in {".m4a", ".aac", ".flac", ".wav", ".ogg", ".opus"}:
        kind = "audio"
    elif mime.startswith("video/") or suffix in {".mkv", ".webm", ".mov", ".avi", ".m4v"}:
        kind = "video"
    elif mime.startswith("text/") or suffix in {".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".yaml", ".yml", ".xml", ".md", ".csv", ".log", ".ini", ".toml", ".sql", ".html", ".css", ".sh", ".bat", ".ps1"}:
        kind = "text"
    elif suffix in {".doc", ".docx"}:
        kind = "document"
    elif suffix in {".xls", ".xlsx", ".xlsm", ".ods"}:
        kind = "spreadsheet"
    elif suffix in {".ppt", ".pptx", ".odp"}:
        kind = "presentation"
    elif suffix == ".pdf":
        kind = "pdf"
    elif suffix in {".zip", ".jar", ".whl", ".epub", ".cbz"} or mime in {"application/zip", "application/x-zip-compressed"}:
        kind = "archive"
    return {"mime_type": mime, "kind": kind, "extension": suffix}


def missing_prerequisites(kind: str) -> list[str]:
    imports = {
        "document": ["docx"],
        "spreadsheet": ["openpyxl"],
        "presentation": ["pptx"],
        "pdf": ["fitz"],
        "image": ["PIL"],
    }
    return [OPTIONAL_PREREQUISITES[name] for name in imports.get(kind, []) if importlib.util.find_spec(name) is None]


def missing_system_prerequisites(kind: str) -> list[str]:
    required = ["ffmpeg", "ffprobe"] if kind in {"audio", "video"} else []
    return [name for name in required if shutil.which(name) is None]


def _system_install_command(package: str) -> list[str] | None:
    if os.name == "nt":
        if shutil.which("winget"):
            return ["winget", "install", "--id", "Gyan.FFmpeg", "-e", "--accept-package-agreements", "--accept-source-agreements"]
        if shutil.which("choco"):
            return ["choco", "install", "ffmpeg", "-y"]
        return None
    if platform.system() == "Darwin" and shutil.which("brew"):
        return ["brew", "install", "ffmpeg"]
    for manager in ("apt-get", "dnf", "pacman"):
        if shutil.which(manager):
            if manager == "apt-get":
                return ["sudo", "apt-get", "install", "-y", "ffmpeg"]
            if manager == "dnf":
                return ["sudo", "dnf", "install", "-y", "ffmpeg"]
            return ["sudo", "pacman", "-S", "--noconfirm", "ffmpeg"]
    return None


def install_known_prerequisites(kind: str, *, install_system: bool = False) -> dict[str, list[str]]:
    assert_mutation_allowed(f"install-prerequisites:{kind}")
    packages = missing_prerequisites(kind)
    installed: list[str] = []
    for package in packages:
        subprocess.run([sys.executable, "-m", "pip", "install", package], check=True, timeout=600)
        installed.append(package)
    system_missing = missing_system_prerequisites(kind)
    system_installed: list[str] = []
    if system_missing and install_system:
        command = _system_install_command("ffmpeg")
        if command is None:
            raise RuntimeError("No supported system package manager was found for FFmpeg.")
        subprocess.run(command, check=True, timeout=900)
        system_installed = [name for name in system_missing if shutil.which(name) is not None]
        if len(system_installed) != len(system_missing):
            raise RuntimeError("FFmpeg installation completed but required binaries are still unavailable.")
    return {"python_installed": installed, "system_missing": system_missing, "system_installed": system_installed}


def extract_text(path: str) -> dict[str, Any]:
    p = Path(path)
    info = detect_type(path)
    if info["kind"] == "text":
        return {"kind": "text", "text": p.read_text(encoding="utf-8", errors="replace")}
    if info["kind"] == "document":
        from docx import Document
        doc = Document(str(p))
        paragraphs = [x.text for x in doc.paragraphs if x.text]
        tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
        return {"kind": "document", "text": "\n".join(paragraphs), "tables": tables}
    if info["kind"] == "spreadsheet":
        from openpyxl import load_workbook
        wb = load_workbook(str(p), read_only=True, data_only=True)
        sheets = {ws.title: [[cell.value for cell in row] for row in ws.iter_rows()] for ws in wb.worksheets}
        return {"kind": "spreadsheet", "sheets": sheets}
    if info["kind"] == "pdf":
        import fitz
        doc = fitz.open(str(p))
        return {"kind": "pdf", "text": "\n".join(page.get_text() for page in doc), "pages": len(doc)}
    raise ValueError(f"No text extractor is registered for {p.suffix or 'this file type'}")


def generic_inspection(path: str, *, max_hash_bytes: int = 16 * 1024 * 1024) -> dict[str, Any]:
    p = Path(path)
    stat = p.stat()
    result: dict[str, Any] = {
        "name": p.name,
        "extension": p.suffix.lower(),
        "size": stat.st_size,
        "modified_at": stat.st_mtime,
        "mime_type": mimetypes.guess_type(p.name)[0] or "application/octet-stream",
    }
    if stat.st_size <= max_hash_bytes:
        digest = hashlib.sha256()
        with p.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        result["sha256"] = digest.hexdigest()
    if detect_type(path)["kind"] == "archive" and zipfile.is_zipfile(p):
        with zipfile.ZipFile(p) as archive:
            result["archive_entries"] = archive.namelist()[:500]
            result["archive_entry_count"] = len(archive.namelist())
    return result


def create_docx(path: str, title: str, paragraphs: list[str]) -> str:
    assert_mutation_allowed("document generation")
    from docx import Document
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(output))
    return str(output.resolve())


def create_xlsx(path: str, sheets: dict[str, list[list[Any]]]) -> str:
    assert_mutation_allowed("spreadsheet generation")
    from openpyxl import Workbook
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31] or "Sheet1"
        for row in rows:
            ws.append(row)
    wb.save(str(output))
    return str(output.resolve())


def create_pdf(path: str, title: str, paragraphs: list[str]) -> str:
    assert_mutation_allowed("pdf generation")
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(output), pagesize=A4)
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    story.extend(Paragraph(p, styles["BodyText"]) for p in paragraphs)
    doc.build(story)
    return str(output.resolve())


def create_pptx(path: str, title: str, slides: list[dict[str, str]]) -> str:
    from pptx import Presentation
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = title
    for slide_data in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(slide_data.get("title", ""))
        slide.placeholders[1].text = str(slide_data.get("body", ""))
    presentation.save(str(output))
    return str(output.resolve())
